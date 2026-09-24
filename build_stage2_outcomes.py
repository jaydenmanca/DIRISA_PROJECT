from collections import defaultdict
from pathlib import Path
import json
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent
DERIVED = ROOT / "derived"
DERIVED.mkdir(exist_ok=True)
BOOKS = [ROOT / "Book1.xlsx", ROOT / "Book2.xlsx"]
BASELINE = ROOT / "LGE_All_Sources_Master_721327_Rows.csv"

def text(v):
    return "" if v is None else str(v).strip()

def num(v):
    try:
        return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0

def read_book(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h:i for i,h in enumerate(header)}
    need = {"Province","Municipality","Ward","VotingDistrict","VotingStationName","RegisteredVoters","BallotType","SpoiltVotes","TotalValidVotes","DateGenerated"}
    missing = sorted(need - set(idx))
    if missing:
        raise ValueError(f"{path.name} missing columns: {missing}")
    units = {}
    provinces, years, municipalities = set(), set(), set()
    data_rows = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        data_rows += 1
        province = text(row[idx["Province"]]); municipality = text(row[idx["Municipality"]])
        ward = text(row[idx["Ward"]]); district = text(row[idx["VotingDistrict"]]); station = text(row[idx["VotingStationName"]])
        registered = num(row[idx["RegisteredVoters"]]); ballot = text(row[idx["BallotType"]])
        spoilt = num(row[idx["SpoiltVotes"]]); valid = num(row[idx["TotalValidVotes"]])
        generated = row[idx["DateGenerated"]]
        year = getattr(generated, "year", None)
        if year is None:
            try: year = pd.to_datetime(generated).year
            except Exception: year = None
        provinces.add(province); municipalities.add(municipality)
        if year is not None and not pd.isna(year): years.add(int(year))
        key = (province, municipality, ward, district, station, registered, ballot, year)
        if key not in units: units[key] = {"valid": 0.0, "spoilt": spoilt}
        units[key]["valid"] += valid
    grouped = defaultdict(lambda: {"registered":0.0,"valid":0.0,"spoilt":0.0,"units":0})
    for (province, municipality, ward, district, station, registered, ballot, year), value in units.items():
        out = grouped[(province, municipality, year, ballot)]
        out["registered"] += registered; out["valid"] += value["valid"]; out["spoilt"] += value["spoilt"]; out["units"] += 1
    rows = []
    for (province, municipality, year, ballot), value in grouped.items():
        cast = value["valid"] + value["spoilt"]
        rows.append({"source_file":path.name,"province":province,"municipality_source":municipality,"election_year":year,"ballot_type":ballot,"registered_voters":value["registered"],"valid_votes":value["valid"],"spoilt_votes":value["spoilt"],"ballots_cast_derived":cast,"turnout_pct_derived":100*cast/value["registered"] if value["registered"] else None,"reporting_units":value["units"],"derivation_note":"valid + spoilt; reconcile with IEC totals"})
    audit = {"source_file":path.name,"sheet":ws.title,"worksheet_rows_including_header":ws.max_row,"data_rows_read":data_rows,"at_excel_row_limit":ws.max_row >= 1048576,"provinces":"; ".join(sorted(provinces)),"election_years":"; ".join(str(x) for x in sorted(years)),"municipality_count":len(municipalities)}
    wb.close()
    return audit, rows

def build_books():
    audits=[]; rows=[]
    for path in BOOKS:
        audit, result = read_book(path); audits.append(audit); rows.extend(result)
    audits.append({"source_file":"Book3.xlsx","sheet":"L3","worksheet_rows_including_header":1048576,"data_rows_read":None,"at_excel_row_limit":True,"provinces":"Eastern Cape (observed in sample)","election_years":"2021 (observed in sample)","municipality_count":None,"status":"blocked_until_original_csv_or_split_source_is_provided"})
    pd.DataFrame(audits).to_csv(DERIVED/"provincial_workbook_audit.csv",index=False)
    pd.DataFrame(rows).sort_values(["province","municipality_source","election_year","ballot_type"]).to_csv(DERIVED/"provincial_book_municipality_election_ballot.csv",index=False)

def build_baseline():
    fields=["election_year","province","municipality_code","municipality_source","ward_code","voting_district","station_name","registered_voters","ballot_type","ballot_votes_cast","spoilt_votes","party_valid_votes_sum","source_filename"]
    units={}
    for chunk in pd.read_csv(BASELINE,usecols=fields,chunksize=150000,low_memory=False):
        for row in chunk.itertuples(index=False,name=None):
            r=dict(zip(fields,row))
            key=tuple(r[x] for x in ["source_filename","election_year","province","municipality_code","municipality_source","ward_code","voting_district","station_name","registered_voters","ballot_type"])
            if key not in units:
                units[key]={x:r[x] for x in ["election_year","province","municipality_code","municipality_source","ballot_type","registered_voters","ballot_votes_cast","spoilt_votes","party_valid_votes_sum"]}
    grouped=defaultdict(lambda: {"registered":0.0,"cast":0.0,"spoilt":0.0,"valid":0.0,"units":0})
    for u in units.values():
        key=(u["province"],u["municipality_code"],u["municipality_source"],u["election_year"],u["ballot_type"])
        out=grouped[key]
        for src,dst in [("registered_voters","registered"),("ballot_votes_cast","cast"),("spoilt_votes","spoilt"),("party_valid_votes_sum","valid")]:
            if pd.notna(u[src]): out[dst]+=float(u[src])
        out["units"]+=1
    rows=[]
    for (province,code,name,year,ballot),v in grouped.items():
        rows.append({"province":province,"municipality_code":code,"municipality_source":name,"election_year":year,"ballot_type":ballot,"registered_voters":v["registered"],"ballot_votes_cast":v["cast"],"spoilt_votes":v["spoilt"],"party_valid_votes_sum":v["valid"],"turnout_pct":100*v["cast"]/v["registered"] if v["registered"] else None,"reporting_units":v["units"],"source_role":"canonical_root_baseline"})
    out=pd.DataFrame(rows).sort_values(["province","municipality_source","election_year","ballot_type"])
    out.to_csv(DERIVED/"national_municipality_election_ballot.csv",index=False)
    return out

if __name__=="__main__":
    build_books()
    out=build_baseline()
    summary={"national_rows":int(len(out)),"national_provinces":int(out["province"].nunique()),"national_municipalities":int(out["municipality_code"].nunique()),"national_years":sorted(int(x) for x in out["election_year"].dropna().unique())}
    (DERIVED/"stage2_summary.json").write_text(json.dumps(summary,indent=2),encoding="utf-8")
    print(json.dumps(summary,indent=2))
